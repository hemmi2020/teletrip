from docx import Document
doc = Document(r"C:/wamp64/www/telitrip/Certification/Pre-Development Questionnaire.docx")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)
for table in doc.tables:
    print("\n--- TABLE ---")
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        print(" | ".join(cells))
